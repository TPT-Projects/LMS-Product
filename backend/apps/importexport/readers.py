"""Read CSV and Excel files into rows of {header: text}.

- CSV: UTF-8 (with or without BOM), falling back to Windows-1252; `,` `;` or tab.
- Excel (.xlsx): first sheet, first non-empty row is the header. Excel stores
  phone numbers as numbers and dates as dates, so they are turned back into text
  (`919876543210`, `05/07/2026`).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from datetime import date, datetime
from typing import IO

from openpyxl import load_workbook

CSV_EXTENSIONS = (".csv", ".txt")
EXCEL_EXTENSIONS = (".xlsx",)


class ReadError(ValueError):
    pass


def file_type_for(name: str) -> str:
    lower = name.lower()
    if lower.endswith(CSV_EXTENSIONS):
        return "csv"
    if lower.endswith(EXCEL_EXTENSIONS):
        return "xlsx"
    raise ReadError("Upload a .csv or .xlsx file.")


def _clean_headers(raw: list) -> list[str]:
    headers, seen = [], {}
    for index, value in enumerate(raw, start=1):
        name = str(value).strip() if value not in (None, "") else f"Column {index}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        headers.append(name if count == 0 else f"{name} ({count + 1})")
    return headers


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute:
            return value.strftime("%d/%m/%Y %H:%M")
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 919876543210.0 -> "919876543210"
    return str(value).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ReadError("Could not read the file. Save it as UTF-8 CSV and try again.")


def _csv_rows(fileobj: IO[bytes]) -> tuple[list[str], Iterator[list[str]]]:
    text = _decode(fileobj.read())
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    for first in reader:
        if any(cell.strip() for cell in first):
            return _clean_headers(first), (row for row in reader)
    raise ReadError("The file is empty.")


def _xlsx_rows(fileobj: IO[bytes]) -> tuple[list[str], Iterator[list[str]]]:
    try:
        workbook = load_workbook(fileobj, read_only=True, data_only=True)
    except Exception as exc:  # corrupt or not really xlsx
        raise ReadError("Could not open the Excel file.") from exc
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    for first in rows:
        if any(v not in (None, "") for v in first):
            headers = _clean_headers(list(first))
            return headers, ([_cell_text(v) for v in row] for row in rows)
    raise ReadError("The first sheet is empty.")


def read_table(fileobj: IO[bytes], file_type: str) -> tuple[list[str], Iterator[dict[str, str]]]:
    """Returns (headers, rows). Rows are dicts and skip lines that are completely empty."""
    headers, raw_rows = _csv_rows(fileobj) if file_type == "csv" else _xlsx_rows(fileobj)

    def rows() -> Iterator[dict[str, str]]:
        for raw in raw_rows:
            values = [str(v).strip() if v is not None else "" for v in raw]
            if not any(values):
                continue
            values += [""] * (len(headers) - len(values))
            yield dict(zip(headers, values, strict=False))

    return headers, rows()
